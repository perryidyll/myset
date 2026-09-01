#!/usr/bin/env python3
"""Build the MySet metrics workbook, for upload to Google Sheets.

    ./metrics.sh && python3 build-sheet.py

Reads whatever CSVs metrics.sh produced in ./metrics/ and writes
metrics/MySet-Metrics.xlsx. Everything derived is a FORMULA, so changing an
assumption re-costs the whole model.
"""
import csv, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

M = "metrics"
FONT = "Arial"
H1   = Font(name=FONT, size=14, bold=True)
HDR  = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, italic=True, color="6E6E73")
INPUT= Font(name=FONT, size=10, color="0000FF")          # hardcoded input
LINK = Font(name=FONT, size=10, color="008000")          # from another sheet
HFILL= PatternFill("solid", fgColor="1D1D1F")
YFILL= PatternFill("solid", fgColor="FFFF00")            # fill this in
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))

wb = Workbook(); wb.remove(wb.active)

def sheet(name, title, blurb, cols=None, widths=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = H1
    ws["A2"] = blurb; ws["A2"].font = NOTE
    if cols:
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=4, column=i, value=c)
            cell.font = HDR; cell.fill = HFILL
            cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.freeze_panes = "A5"
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws

def rows(ws, data, start=5, font=BODY):
    for r, row in enumerate(data, start):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = font; cell.border = THIN
    return start + len(data)

def read(name):
    p = os.path.join(M, name)
    if not os.path.exists(p): return [], []
    with open(p, newline="") as f:
        r = list(csv.reader(f))
    return (r[0], r[1:]) if r else ([], [])

# ── 1 · Read me ───────────────────────────────────────────────────────────
ws = sheet("Read me", "MySet — what we track, and what we don't",
           f"Generated {datetime.date.today().isoformat()} by build-sheet.py. "
           "Re-run ./metrics.sh then build-sheet.py to refresh.")
guide = [
    ["Tab", "What it is", "Where the data comes from", "Refresh"],
    ["Assumptions", "Every lever in the cost model. BLUE cells are the only ones to edit.",
     "Typed by hand; each cell says where the number came from.", "By hand"],
    ["Scale model", "What MySet costs to run at 1 / 10 / 100 / 1000 artists.",
     "Formulas over Assumptions. Nothing here is hardcoded.", "Automatic"],
    ["Infrastructure", "Netlify deploys and credits per site, split by trigger.",
     "metrics.sh -> infra.csv (Netlify API)", "./metrics.sh"],
    ["Gig log", "One row per show: room, voters, votes, songs, money.",
     "metrics.sh -> gigs.csv (needs MYSET_ADMIN_CODE)", "./metrics.sh"],
    ["Song demand", "What the room voted for and NEVER GOT. The best product signal in the app.",
     "metrics.sh -> songs-wanted.csv", "./metrics.sh"],
    ["Song performance", "What actually got played, and what it won with.",
     "metrics.sh -> songs-played.csv", "./metrics.sh"],
    ["Catalogue", "The live song library with genres.",
     "metrics.sh -> catalogue.csv (public API)", "./metrics.sh"],
    ["Not tracked yet", "Honest list of what this sheet CANNOT tell you, and why.",
     "Written by hand.", "By hand"],
]
r = rows(ws, guide[:1], start=4, font=HDR)
for c in range(1, 5): ws.cell(row=4, column=c).fill = HFILL
rows(ws, guide[1:], start=5)
for i, w in enumerate([20, 62, 52, 14], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws["A16"] = "COLOUR KEY"; ws["A16"].font = Font(name=FONT, bold=True, size=10)
ws["A17"] = "Blue text"; ws["A17"].font = INPUT
ws["B17"] = "an input you can change"; ws["B17"].font = BODY
ws["A18"] = "Black text"; ws["A18"].font = BODY
ws["B18"] = "a formula — do not overwrite"; ws["B18"].font = BODY
ws["A19"] = "Yellow fill"; ws["A19"].fill = YFILL; ws["A19"].font = BODY
ws["B19"] = "a key assumption worth arguing about"; ws["B19"].font = BODY

# ── 2 · Assumptions ───────────────────────────────────────────────────────
ws = sheet("Assumptions", "Cost model assumptions",
           "Only the BLUE cells are inputs. Every number cites where it came from.",
           ["Assumption", "Value", "Unit", "Where this number came from"],
           [40, 14, 16, 78])
A = [
 ("Gig length",                    3,      "hours",        "Perry's typical set. Change per artist."),
 ("Phones in the room",            20,     "phones",       "Measured at his 2026-08-30 gig."),
 ("Average seconds between polls", 8,      "seconds",      "vote.html backs off 3s -> 6s -> 12s (INVARIANT 9d). 8s is a conservative mean."),
 ("Gigs per week per artist",      5,      "gigs",         "Perry plays 6-7. Most artists will play far fewer - this is the single biggest lever."),
 ("Payload per poll (gzipped)",    2.5,    "KB",           "MEASURED 2026-09-01: /api/show is 11,035 bytes raw, 2,562 gzipped."),
 ("Non-poll calls per gig",        250,    "calls",        "Votes, payments, admin writes. Estimate."),
 ("Credits per 10,000 requests",   2,      "credits",      "Netlify pricing, recorded in MYSET.md."),
 ("Credits per GB bandwidth",      20,     "credits",      "Netlify pricing."),
 ("Credits per GB-hour compute",   10,     "credits",      "Netlify pricing."),
 ("Function memory",               128,    "MB",           "Netlify default."),
 ("Function duration",             200,    "ms",           "Estimate: /api/show is one strong-consistency read plus 12 shard reads."),
 ("Price per credit",              0.01,   "USD",          "$5 per 500-credit pack. INVARIANT 9d1: never drop to Free, packs are forfeited."),
 ("Revenue per artist",            10.00,  "USD/month",    "The target subscription price."),
 ("Production deploys per month",  20,     "deploys",      "Post-9d3, one per shipped change. Was ~77 when we were deploying twice."),
 ("Credits per production deploy", 15,     "credits",      "INVARIANT 9d0."),
]
for i, (k, v, u, src) in enumerate(A, 5):
    ws.cell(row=i, column=1, value=k).font = BODY
    c = ws.cell(row=i, column=2, value=v); c.font = INPUT
    if k in ("Gigs per week per artist", "Average seconds between polls"): c.fill = YFILL
    if u == "USD": c.number_format = '$#,##0.000'
    if u == "USD/month": c.number_format = '$#,##0.00'
    ws.cell(row=i, column=3, value=u).font = BODY
    ws.cell(row=i, column=4, value=src).font = NOTE
    for col in range(1, 5): ws.cell(row=i, column=col).border = THIN
AS = {k: f"Assumptions!$B${i}" for i, (k, *_ ) in enumerate(A, 5)}

# ── 3 · Scale model ───────────────────────────────────────────────────────
ws = sheet("Scale model", "What MySet costs to run",
           "Every cell is a formula over Assumptions. Change an assumption and this re-costs itself.",
           ["Per artist, per month", "Value", "Unit", "How it is worked out"],
           [38, 16, 14, 76])
g, ph, sec = AS["Gig length"], AS["Phones in the room"], AS["Average seconds between polls"]
gw, kb, nonpoll = AS["Gigs per week per artist"], AS["Payload per poll (gzipped)"], AS["Non-poll calls per gig"]
per = [
 ("Polls per phone per gig",  f"=({g}*3600)/{sec}",                 "polls",    "gig seconds / poll interval"),
 ("Requests per gig",         f"=B5*{ph}+{nonpoll}",                "requests", "polls x phones, plus votes and admin writes"),
 ("Gigs per month",           f"={gw}*52/12",                       "gigs",     "weeks averaged over the year"),
 ("Requests per month",       "=B6*B7",                             "requests", ""),
 ("Bandwidth per month",      f"=B8*{kb}/1048576",                  "GB",       "requests x gzipped payload"),
 ("Compute per month",        f"=B8*{AS['Function duration']}/1000*{AS['Function memory']}/1024/3600", "GB-hours", "duration x memory"),
 ("Credits: requests",        f"=B8/10000*{AS['Credits per 10,000 requests']}", "credits", ""),
 ("Credits: bandwidth",       f"=B9*{AS['Credits per GB bandwidth']}",          "credits", ""),
 ("Credits: compute",         f"=B10*{AS['Credits per GB-hour compute']}",      "credits", ""),
 ("CREDITS PER ARTIST",       "=B11+B12+B13",                       "credits",  "the number that decides whether this scales"),
 ("COST PER ARTIST",          f"=B14*{AS['Price per credit']}",     "USD",      ""),
 ("Gross margin per artist",  f"=({AS['Revenue per artist']}-B15)/{AS['Revenue per artist']}", "%", "at the target subscription price"),
]
for i, (k, f, u, note) in enumerate(per, 5):
    ws.cell(row=i, column=1, value=k).font = Font(name=FONT, size=10, bold=k.isupper())
    c = ws.cell(row=i, column=2, value=f); c.font = BODY
    c.number_format = '$#,##0.00' if u == "USD" else ('0.0%' if u == "%" else '#,##0.00')
    ws.cell(row=i, column=3, value=u).font = BODY
    ws.cell(row=i, column=4, value=note).font = NOTE
    for col in range(1, 5): ws.cell(row=i, column=col).border = THIN

ws["A23"] = "The whole platform"; ws["A23"].font = H1
ws["A24"] = ("Deploys are a FLAT cost — they do not grow with artists. Requests do, "
             "and they are the wall."); ws["A24"].font = NOTE
for i, h in enumerate(["Artists", "Credits/month", "Cost/month", "Revenue/month",
                       "Gross margin", "What breaks first"], 1):
    c = ws.cell(row=26, column=i, value=h); c.font = HDR; c.fill = HFILL
notes = {
 1:    "Nothing. Deploys dominate; that is why 9d3 mattered.",
 10:   "Nothing. Comfortably inside a Personal plan plus packs.",
 100:  "Netlify plan limits. Move to Pro at the END of a cycle (INVARIANT 9d2).",
 1000: "Polling. Not storage - see Not tracked yet. The fix is SSE or a longer backoff, not a new host.",
}
for r_, n in enumerate([1, 10, 100, 1000], 27):
    ws.cell(row=r_, column=1, value=n).font = INPUT
    c = ws.cell(row=r_, column=2, value=f"=$A{r_}*'Scale model'!$B$14+{AS['Production deploys per month']}*{AS['Credits per production deploy']}")
    c.font = BODY; c.number_format = '#,##0'
    c = ws.cell(row=r_, column=3, value=f"=B{r_}*{AS['Price per credit']}"); c.font = BODY; c.number_format = '$#,##0'
    c = ws.cell(row=r_, column=4, value=f"=$A{r_}*{AS['Revenue per artist']}"); c.font = BODY; c.number_format = '$#,##0'
    c = ws.cell(row=r_, column=5, value=f"=IFERROR((D{r_}-C{r_})/D{r_},0)"); c.font = BODY; c.number_format = '0.0%'
    ws.cell(row=r_, column=6, value=notes[n]).font = NOTE
    for col in range(1, 7): ws.cell(row=r_, column=col).border = THIN
for i, w in enumerate([38, 16, 14, 16, 14, 76], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ── 4-8 · the data tabs ───────────────────────────────────────────────────
def data_tab(name, title, blurb, csvname, widths, example=None):
    head, body = read(csvname)
    if not head:
        head = example[0] if example else ["(run ./metrics.sh)"]
    ws = sheet(name, title, blurb, head, widths)
    if body:
        rows(ws, body)
        ws.cell(row=5 + len(body) + 1, column=1,
                value=f"{len(body)} rows from {csvname}, {datetime.date.today().isoformat()}").font = NOTE
    elif example and len(example) > 1:
        rows(ws, example[1:])
        ws.cell(row=6, column=1, value="^ EXAMPLE ROW - shows the expected format. "
                "Delete it once ./metrics.sh has filled this in.").font = NOTE
    return ws

data_tab("Infrastructure", "Netlify: deploys and credits",
         "Split by trigger. via_cli next to via_git on the same site means we are paying twice (INVARIANT 9d3). "
         "Counts the last 100 deploys per site, so it will differ from credit-burn.sh, which counts the billing period.",
         "infra.csv", [26, 42, 24, 20, 12, 12, 12])

data_tab("Gig log", "One row per show",
         "Needs MYSET_ADMIN_CODE. The money columns come from Stripe via /api/history.",
         "gigs.csv", [22, 26, 16, 20, 20, 15, 13, 12, 13, 26, 12, 11, 15],
         example=[["show_id","venue","city","started","ended","phones_in_room","peak_voters",
                   "total_votes","songs_played","top_song","gross_usd","tips_usd","vote_sales_usd"],
                  ["show-1756...","The Ugly Duckling","Koh Phangan","2026-08-30T13:00Z",
                   "2026-08-30T16:00Z",23,14,86,18,"Wonderwall",31.0,25.0,6.0]])

data_tab("Song demand", "What the room wanted and never got",
         "The strongest product signal MySet produces: songs that took votes across a whole night and were never played. "
         "Feeds the songs-to-learn list, and tells you what a crowd in this city actually wants.",
         "songs-wanted.csv", [22, 22, 34, 26, 20],
         example=[["show_id","song_id","title","artist","votes_never_played"],
                  ["show-1756...","mr-brightside","Mr. Brightside","The Killers",11]])

data_tab("Song performance", "What got played, and what it won with",
         "One row per song start. was_replay TRUE means the room paid the higher replay cost to hear it again.",
         "songs-played.csv", [22, 22, 34, 26, 12, 10, 12],
         example=[["show_id","song_id","title","artist","votes_won","voters","was_replay"],
                  ["show-1756...","wonderwall","Wonderwall","Oasis",9,14,False]])

data_tab("Catalogue", "The live song library",
         "Public data - no sign-in needed. votes_now is a snapshot, not a total.",
         "catalogue.csv", [22, 34, 26, 30, 12, 8])

# ── 9 · Not tracked yet ───────────────────────────────────────────────────
ws = sheet("Not tracked yet", "What this sheet cannot tell you",
           "Written by hand, deliberately. A metrics sheet that hides its gaps is worse than none.",
           ["Question you will want to ask", "Why it is not here", "What it would take"],
           [50, 62, 56])
gaps = [
 ["How many people scanned the QR and never voted?",
  "MySet has NO analytics. Not one page view is recorded anywhere - verified by grep across every page and function.",
  "A single counter on /api/show when in=1 is sent. Presence is already stamped per device (INVARIANT 0af); it just is not counted for people who never vote."],
 ["What is the conversion from room to paying?",
  "Presence is counted (phones in the room) and payments are recorded, but nothing joins them into a funnel.",
  "Join gigs.csv phones_in_room against gross_usd - possible TODAY in this sheet, just not automatic."],
 ["Which marketing channel brought an artist?",
  "?ref= is recorded at signup for referrals only. utm_* params are STRIPPED (_profile.mjs, _venues.mjs).",
  "Keep utm_source at signup on the artist record. Small change, real answer."],
 ["How many artists are there, and what plan?",
  "The registry has it, but no endpoint exposes a roll-up - and rightly, since it is not any one artist's business.",
  "A platform-owner-only endpoint gated on isPlatformOwner (the same gate promo codes use)."],
 ["Why did an artist stop using it?",
  "No retention or churn data of any kind.",
  "Last-gig date per artist is derivable from history today. Churn needs that plus a definition."],
 ["Does the audience come back?",
  "Fan ids are per-device and per-artist, and are deliberately anonymous (INVARIANT 9g).",
  "Genuinely hard, and worth NOT solving. Anonymity is why the app works in a bar."],
 ["What does storage actually cost?",
  "Not measured, because it is not close to mattering. A 65-song show doc is ~20KB; a whole artist-year including history is a few MB.",
  "Nothing. Revisit if artists ever exceed ~10,000."],
]
rows(ws, gaps)

os.makedirs(M, exist_ok=True)
out = os.path.join(M, "MySet-Metrics.xlsx")
wb.save(out)
print("wrote", out)
